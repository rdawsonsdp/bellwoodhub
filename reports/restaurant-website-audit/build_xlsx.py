import json, glob, os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

SP = "/tmp/claude-0/-home-user-bellwoodhub/c81ee16d-be56-5745-928f-8e25e1544da1/scratchpad"

zones = []
for f in sorted(glob.glob(os.path.join(SP, "zone*.json"))):
    with open(f) as fh:
        zones.append(json.load(fh))

STATUS_LABEL = {
    "NO_WEBSITE": "No website",
    "WEBSITE_UNAVAILABLE": "Website unavailable",
    "SOCIAL_ONLY": "Social media only",
}
STATUS_ORDER = {"NO_WEBSITE": 0, "WEBSITE_UNAVAILABLE": 1, "SOCIAL_ONLY": 2}

rows = []
for z in zones:
    for r in z["flagged"]:
        rows.append({
            "name": r.get("name", ""),
            "address": r.get("address", ""),
            "town": r.get("town", ""),
            "phone": r.get("phone", ""),
            "cuisine": r.get("cuisine", ""),
            "status": r.get("status", ""),
            "url": r.get("url_found", ""),
            "notes": r.get("notes", ""),
            "zone": z["zone"],
        })

# dedupe by normalized name+town
seen = set()
deduped = []
for r in rows:
    key = ("".join(c for c in r["name"].lower() if c.isalnum()),
           "".join(c for c in r["town"].lower() if c.isalpha()))
    if key in seen:
        continue
    seen.add(key)
    deduped.append(r)
rows = sorted(deduped, key=lambda r: (STATUS_ORDER.get(r["status"], 9), r["town"], r["name"]))

checked = sum(z["checked"] for z in zones)
has_site = sum(z["has_website"] for z in zones)

wb = Workbook()

NAVY = "1F3864"
CREAM = "FBF6EE"
thin = Side(style="thin", color="B0B0B0")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
hdr_font = Font(name="Arial", bold=True, color="FFFFFF", size=10)
hdr_fill = PatternFill("solid", fgColor=NAVY)
body = Font(name="Arial", size=10)

# ---- Sheet 1: Flagged Restaurants ----
ws = wb.active
ws.title = "Flagged Restaurants"
headers = ["#", "Restaurant", "Address", "Town / Neighborhood", "Phone", "Cuisine",
           "Status", "URL Found (broken or social)", "Evidence Notes"]
ws.append(headers)
for c, h in enumerate(headers, 1):
    cell = ws.cell(row=1, column=c)
    cell.font = hdr_font
    cell.fill = hdr_fill
    cell.alignment = Alignment(vertical="center", wrap_text=True)
    cell.border = border

status_fill = {
    "NO_WEBSITE": PatternFill("solid", fgColor="FDE9E9"),
    "WEBSITE_UNAVAILABLE": PatternFill("solid", fgColor="FFF3DE"),
    "SOCIAL_ONLY": PatternFill("solid", fgColor="EAF1FB"),
}

for i, r in enumerate(rows, 1):
    ws.append([i, r["name"], r["address"], r["town"], r["phone"], r["cuisine"],
               STATUS_LABEL.get(r["status"], r["status"]), r["url"], r["notes"]])
    for c in range(1, 10):
        cell = ws.cell(row=i + 1, column=c)
        cell.font = body
        cell.border = border
        cell.alignment = Alignment(vertical="top", wrap_text=(c in (2, 3, 6, 8, 9)))
    ws.cell(row=i + 1, column=7).fill = status_fill.get(r["status"], PatternFill())

widths = [4, 26, 30, 22, 14, 26, 18, 40, 60]
for c, w in enumerate(widths, 1):
    ws.column_dimensions[get_column_letter(c)].width = w
ws.freeze_panes = "A2"
ws.auto_filter.ref = f"A1:I{len(rows) + 1}"

# ---- Sheet 2: Summary ----
n = len(rows)
s = wb.create_sheet("Summary")
s["A1"] = "Restaurant Website Audit — 25-Mile Radius of Bellwood, IL"
s["A1"].font = Font(name="Arial", bold=True, size=14, color=NAVY)
s["A2"] = "Method: web search per restaurant for an official website; each candidate URL fetched and verified. Independent/local restaurants only (national chains excluded). Best-effort sweep, not a census."
s["A2"].font = Font(name="Arial", size=9, italic=True)
s["A2"].alignment = Alignment(wrap_text=True)
s.merge_cells("A2:D2")
s.row_dimensions[2].height = 30
s["A3"] = "Run date: July 22, 2026"
s["A3"].font = Font(name="Arial", size=9, italic=True)

by_status = {}
for r in rows:
    by_status[r["status"]] = by_status.get(r["status"], 0) + 1
items = [
    ("Restaurants checked", checked, None),
    ("Had a working website", has_site, None),
    ("Flagged (no/broken/social-only site)", n, None),
    ("— No website at all", by_status.get("NO_WEBSITE", 0), None),
    ("— Website unavailable (dead/parked/broken)", by_status.get("WEBSITE_UNAVAILABLE", 0), None),
    ("— Social media / delivery apps only", by_status.get("SOCIAL_ONLY", 0), None),
]
row = 5
s.cell(row=row, column=1, value="Metric").font = hdr_font
s.cell(row=row, column=1).fill = hdr_fill
s.cell(row=row, column=2, value="Count").font = hdr_font
s.cell(row=row, column=2).fill = hdr_fill
for label, val, match in items:
    row += 1
    s.cell(row=row, column=1, value=label).font = body
    c = s.cell(row=row, column=2)
    c.value = val
    c.font = body
row += 2
s.cell(row=row, column=1, value="Flagged by zone").font = Font(name="Arial", bold=True, size=11, color=NAVY)
row += 1
s.cell(row=row, column=1, value="Zone").font = hdr_font
s.cell(row=row, column=1).fill = hdr_fill
s.cell(row=row, column=2, value="Checked").font = hdr_font
s.cell(row=row, column=2).fill = hdr_fill
s.cell(row=row, column=3, value="Flagged").font = hdr_font
s.cell(row=row, column=3).fill = hdr_fill
for z in zones:
    row += 1
    s.cell(row=row, column=1, value=z["zone"]).font = body
    s.cell(row=row, column=2, value=z["checked"]).font = body
    s.cell(row=row, column=3, value=len(z["flagged"])).font = body
s.column_dimensions["A"].width = 58
s.column_dimensions["B"].width = 10
s.column_dimensions["C"].width = 10

out = os.path.join(SP, "restaurants-no-website-bellwood-25mi.xlsx")
wb.save(out)
print("rows:", n, "checked:", checked, "has_site:", has_site)
print("saved:", out)

# also dump merged JSON for the PDF step
with open(os.path.join(SP, "merged.json"), "w") as fh:
    json.dump({"rows": rows, "checked": checked, "has_site": has_site,
               "zones": [{"zone": z["zone"], "checked": z["checked"], "flagged": len(z["flagged"])} for z in zones]}, fh, indent=1)
